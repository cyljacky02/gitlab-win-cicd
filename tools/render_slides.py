"""Build a concise Markdown deck with native editable flowcharts.

Inputs: # title, > takeaway, ## cards, - bullets, Mermaid flowchart LR,
and HTML speaker notes. A standalone ![alt](relative.png) creates a full-slide
image page; its editable source is maintained separately. Mermaid comments @node/@lane/@route/@label carry
explicit slide coordinates; Mermaid itself ignores them. This is a strict
subset renderer, not a general-purpose Mermaid layout engine.
"""
from pathlib import Path
import argparse
import re
import os
import math
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

W, H = 960, 540
NAVY, TEAL, INK = '#142C42', '#087F8C', '#23384D'
MUTED, BG, LINE, ORANGE = '#52677B', '#F3F6FA', '#DCE5EE', '#C57227'
FONT = 'Microsoft JhengHei'
PALETTE = {
    'default': (NAVY, '#FFFFFF', NAVY),
    'target': (TEAL, '#FFFFFF', TEAL),
    'control': ('#FFF1DA', '#754519', ORANGE),
    'blocked': ('#FCE8E6', '#922A2A', '#A63030'),
}


def parse_mermaid(source):
    graph = dict(source=source.strip(), nodes={}, lanes={}, edges=[], routes={}, labels={}, ports={}, name=None)
    lane = None
    for raw in source.strip().splitlines():
        line = raw.strip()
        if not line:
            continue
        if line == 'flowchart LR':
            continue
        m = re.fullmatch(r'%% @name ([a-z0-9-]+)', line)
        if m:
            graph['name'] = m[1]
            continue
        m = re.fullmatch(r'%% @(node|lane) (\w+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)', line)
        if m:
            collection = graph['nodes'] if m[1] == 'node' else graph['lanes']
            collection.setdefault(m[2], {}).update(zip(('x', 'y', 'w', 'h'), map(float, m.groups()[2:])))
            continue
        m = re.fullmatch(r'%% @ports (\w+) (\w+) (left|right|top|bottom) (left|right|top|bottom)', line)
        if m:
            graph['ports'][(m[1], m[2])] = (m[3], m[4])
            continue
        m = re.fullmatch(r'%% @route (\w+) (\w+) (.+)', line)
        if m:
            values = list(map(float, m[3].split()))
            if len(values) % 2:
                raise ValueError('Route needs coordinate pairs')
            graph['routes'][(m[1], m[2])] = list(zip(values[::2], values[1::2]))
            continue
        m = re.fullmatch(r'%% @label (\w+) (\w+) ([\d.]+) ([\d.]+) ([\d.]+)', line)
        if m:
            graph['labels'][(m[1], m[2])] = tuple(map(float, m.groups()[2:]))
            continue
        m = re.fullmatch(r'subgraph (\w+)\["(.*?)"\]', line)
        if m:
            lane = m[1]
            graph['lanes'].setdefault(lane, {})['label'] = m[2]
            continue
        if line == 'end':
            lane = None
            continue
        m = re.fullmatch(r'(\w+)(\[|\{)"(.*?)"(\]|\})', line)
        if m:
            if (m[2], m[4]) not in [('[', ']'), ('{', '}')]:
                raise ValueError('Mismatched node delimiters')
            graph['nodes'].setdefault(m[1], {}).update(
                label=m[3].replace('<br/>', '\n'), lane=lane,
                shape='diamond' if m[2] == '{' else 'round', category='default')
            continue
        m = re.fullmatch(r'(\w+)\s+(-->|-\.->)(?:\|"(.*?)"\|)?\s+(\w+)', line)
        if m:
            graph['edges'].append(dict(source=m[1], target=m[4], label=m[3] or '', dashed=m[2] == '-.->'))
            continue
        m = re.fullmatch(r'class ([\w,]+) (\w+)', line)
        if m:
            if m[2] not in PALETTE:
                raise ValueError(f'Unsupported category: {m[2]}')
            for node in m[1].split(','):
                graph['nodes'][node]['category'] = m[2]
            continue
        if line.startswith('classDef '):
            # Validate class styling rather than silently ignoring arbitrary CSS.
            m = re.fullmatch(r'classDef (\w+) fill:(#[0-9A-Fa-f]{6}),color:(#[0-9A-Fa-f]{6}),stroke:(#[0-9A-Fa-f]{6})', line)
            if not m or m[1] not in PALETTE:
                raise ValueError(f'Unsupported Mermaid class: {line}')
            continue
        raise ValueError(f'Unsupported Mermaid syntax: {line}')
    if not graph['name'] or not graph['nodes']:
        raise ValueError('Diagram needs a @name and nodes')
    for group in ('nodes', 'lanes'):
        for key, value in graph[group].items():
            if not all(k in value for k in ('x', 'y', 'w', 'h', 'label')):
                raise ValueError(f'Missing label/geometry: {key}')
            if value['x'] < 0 or value['y'] < 0 or value['x'] + value['w'] > 880 or value['y'] + value['h'] > 280:
                raise ValueError(f'Diagram geometry outside 880x280: {key}')
            if group == 'nodes' and value['lane']:
                parent = graph['lanes'][value['lane']]
                if not (parent['x'] <= value['x'] and parent['y'] + 32 <= value['y']
                        and value['x'] + value['w'] <= parent['x'] + parent['w']
                        and value['y'] + value['h'] <= parent['y'] + parent['h']):
                    raise ValueError(f'Node {key} outside its responsibility lane')
    for edge in graph['edges']:
        if edge['source'] not in graph['nodes'] or edge['target'] not in graph['nodes']:
            raise ValueError('Edge references undefined node')
    return graph


def parse(source):
    slides = []
    for block in re.split(r'^---\s*$', source, flags=re.M):
        notes = re.findall(r'<!-- notes\s*(.*?)-->', block, re.S)
        block = re.sub(r'<!-- notes\s*.*?-->', '', block, flags=re.S)
        diagrams = re.findall(r'```mermaid\s*\n(.*?)\n```', block, re.S)
        block = re.sub(r'```mermaid\s*\n.*?\n```', '', block, flags=re.S)
        images = re.findall(r'^!\[([^\]]+)\]\(([^)]+)\)\s*$', block, flags=re.M)
        block = re.sub(r'^!\[[^\]]+\]\([^)]+\)\s*$', '', block, flags=re.M)
        title, takeaway, cards = None, None, []
        for line in block.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith('# '):
                if title:
                    raise ValueError('Multiple slide titles')
                title = line[2:]
            elif line.startswith('> '):
                takeaway = line[2:]
            elif line.startswith('## '):
                cards.append({'title': line[3:], 'bullets': []})
            elif line.startswith('- ') and cards:
                cards[-1]['bullets'].append(line[2:])
            else:
                raise ValueError(f'Unsupported Markdown: {line}')
        if not title or not takeaway or len(diagrams) > 1 or len(images) > 1:
            raise ValueError('Slide needs title, takeaway, <=1 diagram/image')
        if images:
            if diagrams or cards:
                raise ValueError('Full-slide image cannot contain cards or Mermaid')
        elif not 1 <= len(cards) <= 3:
            raise ValueError('Standard slide needs 1–3 cards')
        if diagrams and (len(cards) != 1 or len(cards[0]['bullets']) != 1):
            raise ValueError('Diagram slide must have one short explanatory sentence')
        slides.append(dict(title=title, takeaway=takeaway, cards=cards,
                           diagram=parse_mermaid(diagrams[0]) if diagrams else None,
                           image=dict(alt=images[0][0], path=images[0][1]) if images else None,
                           notes='\n\n'.join(notes).strip()))
    return slides


def register_fonts(directory):
    for alias, name in [('regular', 'msjh.ttc'), ('bold', 'msjhbd.ttc')]:
        file = directory / name
        if not file.exists():
            raise FileNotFoundError(f'Chinese font missing: {file}; use --font-dir')
        pdfmetrics.registerFont(TTFont(alias, str(file), subfontIndex=0))


def wrapped(value, width, size, bold=False):
    alias = 'bold' if bold else 'regular'
    result = []
    for paragraph in value.split('\n'):
        line = ''
        tokens = re.findall(r'[A-Za-z0-9_./?=:+\-]+|\s+|.', paragraph)
        for token in tokens:
            pieces = list(token) if pdfmetrics.stringWidth(token, alias, size) > width else [token]
            for piece in pieces:
                if line and pdfmetrics.stringWidth(line + piece, alias, size) > width:
                    result.append(line.rstrip())
                    line = piece.lstrip()
                else:
                    line += piece
        result.append(line.rstrip())
    return result


def text_height(value, width, size, bold=False):
    return len(wrapped(value, width, size, bold)) * size * 1.35


def shape(scene, x, y, w, h, fill, stroke=None, form='rect', label=''):
    scene.append(dict(kind='shape', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, form=form, label=label))


def text(scene, value, x, y, width, size=20, color=INK, bold=False, align='left', role=None):
    lines = wrapped(value, width, size, bold)
    height = len(lines) * size * 1.35
    if x < 0 or y < 0 or x + width > W + .1 or y + height > H + .1:
        raise ValueError(f'Text outside slide: {value}')
    scene.append(dict(kind='text', lines=lines, x=x, y=y, w=width,
                      size=size, color=color, bold=bold, align=align, role=role))
    return height


def edge_points(graph, edge):
    a, b = graph['nodes'][edge['source']], graph['nodes'][edge['target']]
    ac, bc = (a['x'] + a['w']/2, a['y'] + a['h']/2), (b['x'] + b['w']/2, b['y'] + b['h']/2)
    if b['x'] >= a['x'] + a['w']:
        start, end = (a['x'] + a['w'], ac[1]), (b['x'], bc[1])
    elif a['x'] >= b['x'] + b['w']:
        start, end = (a['x'], ac[1]), (b['x'] + b['w'], bc[1])
    elif bc[1] > ac[1]:
        common_x = min(max(ac[0], max(a['x'], b['x'])), min(a['x'] + a['w'], b['x'] + b['w']))
        start, end = (common_x, a['y'] + a['h']), (common_x, b['y'])
    else:
        common_x = min(max(ac[0], max(a['x'], b['x'])), min(a['x'] + a['w'], b['x'] + b['w']))
        start, end = (common_x, a['y']), (common_x, b['y'] + b['h'])
    ports = graph['ports'].get((edge['source'], edge['target']))
    if ports:
        def anchor(node, side):
            return {
                'left': (node['x'], node['y'] + node['h']/2),
                'right': (node['x'] + node['w'], node['y'] + node['h']/2),
                'top': (node['x'] + node['w']/2, node['y']),
                'bottom': (node['x'] + node['w']/2, node['y'] + node['h']),
            }[side]
        start, end = anchor(a, ports[0]), anchor(b, ports[1])
    points = graph['routes'].get((edge['source'], edge['target']), [])
    if not points and start[0] != end[0] and start[1] != end[1]:
        middle = (start[0] + end[0])/2
        points = [(middle, start[1]), (middle, end[1])]
    return [start, *points, end]


def diagram_scene(scene, graph):
    ox, oy = 40, 165
    for lane in graph['lanes'].values():
        shape(scene, ox + lane['x'], oy + lane['y'], lane['w'], lane['h'], '#FFFFFF', LINE, 'round')
        shape(scene, ox + lane['x'], oy + lane['y'], lane['w'], 34, '#E4EBF3')
        text(scene, lane['label'], ox + lane['x'] + 12, oy + lane['y'] + 4, lane['w'] - 24, 18, NAVY, True)
    for edge in graph['edges']:
        points = edge_points(graph, edge)
        scene.append(dict(kind='arrow', points=[(ox+x, oy+y) for x, y in points],
                          color='#A63030' if edge['dashed'] else TEAL, dashed=edge['dashed']))
    for node in graph['nodes'].values():
        fill, color, stroke = PALETTE[node['category']]
        x, y, w, h = node['x'] + ox, node['y'] + oy, node['w'], node['h']
        shape(scene, x, y, w, h, fill, stroke, node['shape'], node['label'])
        size = 20
        width = w * .70 if node['shape'] == 'diamond' else w - 22
        max_height = h * .63 if node['shape'] == 'diamond' else h - 14
        while text_height(node['label'], width, size, True) > max_height and size > 18:
            size -= .5
        height = text_height(node['label'], width, size, True)
        if height > max_height:
            raise ValueError(f'Diagram label too long: {node["label"]}')
        text(scene, node['label'], x + (w-width)/2, y + (h-height)/2 - 1, width, size, color, True, 'center')
    for edge in graph['edges']:
        if not edge['label']:
            continue
        custom = graph['labels'].get((edge['source'], edge['target']))
        if custom:
            x, y, width = custom
        else:
            points = edge_points(graph, edge)
            longest = max(zip(points, points[1:]), key=lambda pair: math.dist(*pair))
            cx, cy = (longest[0][0]+longest[1][0])/2, (longest[0][1]+longest[1][1])/2
            width = max(80, pdfmetrics.stringWidth(edge['label'], 'bold', 18) + 16)
            x, y = cx-width/2, cy-32
        color = '#922A2A' if edge['dashed'] else TEAL
        shape(scene, ox+x, oy+y, width, 27, BG)
        text(scene, edge['label'], ox+x, oy+y, width, 18, color, True, 'center')


def make_scene(slide, index, count):
    if slide.get('image'):
        from reportlab.lib.utils import ImageReader
        image = slide['image']
        iw, ih = ImageReader(str(image['path'])).getSize()
        scale = min(W / iw, H / ih)
        return [dict(kind='image', path=str(image['path']), alt=image['alt'],
                     title=slide['title'], x=(W-iw*scale)/2, y=(H-ih*scale)/2,
                     w=iw*scale, h=ih*scale)]
    scene = []
    cover = index == 0
    shape(scene, 0, 0, W, H, NAVY if cover else BG)
    shape(scene, 0, 0, 9, H, TEAL)
    text(scene, 'WINDOWS DEPLOYMENT  /  跨部門討論提案', 40, 25, 880, 12,
         '#A8CCD6' if cover else TEAL, True)
    if cover:
        shape(scene, 40, 114, 68, 5, '#40BAC5')
        text(scene, slide['title'], 40, 145, 870, 44, '#FFFFFF', True, role='title')
        text(scene, slide['takeaway'], 40, 235, 850, 24, '#D3E4EB')
        text(scene, slide['cards'][0]['title'], 40, 340, 850, 23, '#FFFFFF', True)
        for i, bullet in enumerate(slide['cards'][0]['bullets']):
            text(scene, bullet, 40, 385 + i*35, 850, 18, '#D3E4EB')
    else:
        size = 32
        while text_height(slide['title'], 880, size, True) > 46 and size > 26:
            size -= .5
        text(scene, slide['title'], 40, 62, 880, size, NAVY, True, role='title')
        text(scene, slide['takeaway'], 40, 113, 880, 21, MUTED)
        if slide['diagram']:
            diagram_scene(scene, slide['diagram'])
            text(scene, slide['cards'][0]['bullets'][0], 40, 462, 880, 20, INK)
        else:
            top, bottom = 190, 480
            n = len(slide['cards'])
            cw = (880 - 18*(n-1))/n
            for i, card in enumerate(slide['cards']):
                x = 40 + i*(cw+18)
                shape(scene, x, top, cw, bottom-top, '#FFFFFF', LINE, 'round')
                shape(scene, x, top, cw, 4, TEAL if i < n-1 else ORANGE)
                head_h = text(scene, card['title'], x+16, top+18, cw-32, 21, NAVY, True)
                body_y = top + 18 + head_h + 20
                available = bottom - 16 - body_y
                size = 22
                def required(s):
                    return sum(text_height(b, cw-46, s) + 16 for b in card['bullets']) - 16
                while required(size) > available and size > 18:
                    size -= .5
                if required(size) > available:
                    raise ValueError(f'Card overflow on slide {index+1}: {card["title"]}')
                for bullet in card['bullets']:
                    text(scene, '•', x+15, body_y, 12, size, TEAL, True)
                    body_y += text(scene, bullet, x+30, body_y, cw-46, size) + 16
    shape(scene, 40, 505, 880, 1, '#365164' if cover else LINE)
    text(scene, 'v0.7  |  初步計劃 · 正式上線另需批准', 40, 514, 760, 10,
         '#A8CCD6' if cover else MUTED)
    text(scene, f'{index+1:02d} / {count:02d}', 859, 514, 61, 10,
         '#A8CCD6' if cover else MUTED)
    return scene


def rgb(value):
    return RGBColor.from_string(value.lstrip('#'))


def ppt_scene(prs, scene, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for item in scene:
        if item['kind'] == 'image':
            # Preserve a semantic slide title behind the full-page image.
            slide.shapes.title.text = item['title']
            background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Pt(W), Pt(H))
            background.fill.solid()
            background.fill.fore_color.rgb = rgb('#FFFFFF')
            background.line.fill.background()
            picture = slide.shapes.add_picture(item['path'], Pt(item['x']), Pt(item['y']),
                                               width=Pt(item['w']), height=Pt(item['h']))
            picture._element.nvPicPr.cNvPr.set('descr', item['alt'])
            continue
        if item['kind'] == 'arrow':
            pairs = list(zip(item['points'], item['points'][1:]))
            for i, ((x1, y1), (x2, y2)) in enumerate(pairs):
                arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
                arrow.line.color.rgb = rgb(item['color'])
                arrow.line.width = Pt(2)
                if item['dashed']:
                    arrow.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                if i == len(pairs)-1:
                    head = OxmlElement('a:tailEnd')
                    head.set('type', 'triangle')
                    head.set('w', 'sm')
                    head.set('len', 'sm')
                    arrow.line._get_or_add_ln().append(head)
            continue
        x, y, w = item['x'], item['y'], item['w']
        if item['kind'] == 'shape':
            form = {'rect': MSO_SHAPE.RECTANGLE, 'round': MSO_SHAPE.ROUNDED_RECTANGLE,
                    'diamond': MSO_SHAPE.DIAMOND}[item['form']]
            s = slide.shapes.add_shape(form, Pt(x), Pt(y), Pt(w), Pt(item['h']))
            if item['form'] == 'round':
                s.adjustments[0] = .08
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(item['fill'])
            if item['stroke']:
                s.line.color.rgb = rgb(item['stroke'])
                s.line.width = Pt(.8)
            else:
                s.line.fill.background()
            if item['label']:
                s._element.nvSpPr.cNvPr.set('descr', item['label'])
        else:
            size = item['size']
            if item['role'] == 'title':
                box = slide.shapes.title
                box.left, box.top, box.width, box.height = Pt(x), Pt(y), Pt(w), Pt(len(item['lines'])*size*1.35+2)
                # Move the title above the earlier background shapes.
                parent = box._element.getparent()
                parent.remove(box._element)
                parent.append(box._element)
            else:
                box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(len(item['lines'])*size*1.35+2))
            tf = box.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for i, line in enumerate(item['lines']):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.alignment = PP_ALIGN.CENTER if item['align'] == 'center' else PP_ALIGN.LEFT
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = Pt(size*1.35)
                p.font.name = FONT
                p.font.size = Pt(size)
                p.font.bold = item['bold']
                p.font.color.rgb = rgb(item['color'])
    slide.notes_slide.notes_text_frame.text = notes


def pdf_scene(canvas, scene):
    for item in scene:
        if item['kind'] == 'image':
            canvas.drawImage(item['path'], item['x'], H-item['y']-item['h'],
                             width=item['w'], height=item['h'])
        elif item['kind'] == 'arrow':
            canvas.setStrokeColor(item['color'])
            canvas.setFillColor(item['color'])
            canvas.setLineWidth(2)
            canvas.setDash(5, 4) if item['dashed'] else canvas.setDash()
            for (x1, y1), (x2, y2) in zip(item['points'], item['points'][1:]):
                canvas.line(x1, H-y1, x2, H-y2)
            a, b = item['points'][-2:]
            angle = math.atan2(b[1]-a[1], b[0]-a[0])
            path = canvas.beginPath()
            path.moveTo(b[0], H-b[1])
            for side in (-1, 1):
                theta = angle + side*math.pi/6
                path.lineTo(b[0]-9*math.cos(theta), H-(b[1]-9*math.sin(theta)))
            path.close()
            canvas.setDash()
            canvas.drawPath(path, fill=1, stroke=0)
        elif item['kind'] == 'shape':
            canvas.setFillColor(item['fill'])
            canvas.setStrokeColor(item['stroke'] or item['fill'])
            canvas.setLineWidth(.8)
            x, y, w, h = item['x'], H-item['y']-item['h'], item['w'], item['h']
            stroke = int(bool(item['stroke']))
            if item['form'] == 'diamond':
                path = canvas.beginPath()
                path.moveTo(x+w/2, y+h)
                path.lineTo(x+w, y+h/2)
                path.lineTo(x+w/2, y)
                path.lineTo(x, y+h/2)
                path.close()
                canvas.drawPath(path, fill=1, stroke=stroke)
            elif item['form'] == 'round':
                canvas.roundRect(x, y, w, h, min(w,h)*.08, fill=1, stroke=stroke)
            else:
                canvas.rect(x, y, w, h, fill=1, stroke=stroke)
        else:
            canvas.setFillColor(item['color'])
            alias = 'bold' if item['bold'] else 'regular'
            canvas.setFont(alias, item['size'])
            for i, line in enumerate(item['lines']):
                x = item['x']
                if item['align'] == 'center':
                    x += (item['w']-pdfmetrics.stringWidth(line, alias, item['size']))/2
                canvas.drawString(x, H-item['y']-item['size']*(1.08+i*1.35), line)
    canvas.showPage()


def export_diagram(graph, directory):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / (graph['name']+'.mmd')).write_text(graph['source']+'\n', encoding='utf-8')
    file = ET.Element('mxfile', host='app.diagrams.net', type='device')
    diagram = ET.SubElement(file, 'diagram', name=graph['name'], id=graph['name'])
    model = ET.SubElement(diagram, 'mxGraphModel', grid='1', page='1', pageWidth='880', pageHeight='280')
    root = ET.SubElement(model, 'root')
    ET.SubElement(root, 'mxCell', id='0')
    ET.SubElement(root, 'mxCell', id='1', parent='0')
    def vertex(key, value, parent, style, x=None, y=None):
        cell = ET.SubElement(root, 'mxCell', id=key, value=value['label'].replace('\n','<br/>'),
                             vertex='1', parent=parent, style=style)
        ET.SubElement(cell, 'mxGeometry', x=str(value['x'] if x is None else x),
                      y=str(value['y'] if y is None else y), width=str(value['w']), height=str(value['h']), **{'as':'geometry'})
    for key, lane in graph['lanes'].items():
        vertex(key, lane, '1', 'swimlane;horizontal=1;startSize=34;fillColor=#E4EBF3;swimlaneFillColor=#FFFFFF;fontSize=18;fontStyle=1;strokeColor=#DCE5EE;')
    for key, node in graph['nodes'].items():
        fill, color, stroke = PALETTE[node['category']]
        style = ('rhombus;' if node['shape']=='diamond' else 'rounded=1;arcSize=8;')
        style += f'html=1;whiteSpace=wrap;fontSize=20;fontStyle=1;fillColor={fill};fontColor={color};strokeColor={stroke};'
        parent = node['lane'] or '1'
        lane = graph['lanes'].get(parent, {'x':0,'y':0})
        vertex(key, node, parent, style, node['x']-lane['x'], node['y']-lane['y'])
    for i, edge in enumerate(graph['edges']):
        color = '#A63030' if edge['dashed'] else TEAL
        style = f'edgeStyle=orthogonalEdgeStyle;rounded=0;endArrow=block;endFill=1;strokeWidth=2;fontSize=18;strokeColor={color};'
        if edge['dashed']:
            style += 'dashed=1;'
        ports = graph['ports'].get((edge['source'], edge['target']))
        if ports:
            sides = {'left': (0, .5), 'right': (1, .5), 'top': (.5, 0), 'bottom': (.5, 1)}
            exit_x, exit_y = sides[ports[0]]
            entry_x, entry_y = sides[ports[1]]
            style += f'exitX={exit_x};exitY={exit_y};entryX={entry_x};entryY={entry_y};exitPerimeter=1;entryPerimeter=1;'
        cell = ET.SubElement(root, 'mxCell', id=f'edge-{i}', value=edge['label'], edge='1', parent='1',
                             source=edge['source'], target=edge['target'], style=style)
        geometry = ET.SubElement(cell, 'mxGeometry', relative='1', **{'as':'geometry'})
        points = graph['routes'].get((edge['source'],edge['target']), [])
        if points:
            array = ET.SubElement(geometry, 'Array', **{'as':'points'})
            for x,y in points:
                ET.SubElement(array, 'mxPoint', x=str(x), y=str(y))
    ET.indent(file)
    ET.ElementTree(file).write(directory/(graph['name']+'.drawio'), encoding='utf-8', xml_declaration=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--font-dir', type=Path, default=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts')
    args = parser.parse_args()
    register_fonts(args.font_dir)
    slides = parse(args.source.read_text(encoding='utf-8'))
    for slide in slides:
        if slide['image']:
            image_path = args.source.parent / slide['image']['path']
            if not image_path.is_file():
                raise FileNotFoundError(f'Slide image missing: {image_path}')
            slide['image']['path'] = image_path
    output = args.output_dir or args.source.parent
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    prs.core_properties.title = 'Windows 正式部署治理計劃'
    prs.core_properties.subject = '跨部門可行性與合規討論；非正式上線核准'
    prs.core_properties.author = 'Windows CD Governance Project'
    pdf_path, ppt_path = [output/args.source.with_suffix('.'+ext).name for ext in ('pdf','pptx')]
    canvas = Canvas(str(pdf_path), pagesize=(W,H), pageCompression=1)
    canvas.setTitle(prs.core_properties.title)
    canvas.setAuthor(prs.core_properties.author)
    scenes = [make_scene(s,i,len(slides)) for i,s in enumerate(slides)]
    for slide, scene in zip(slides, scenes):
        ppt_scene(prs, scene, slide['notes'])
        pdf_scene(canvas, scene)
        if slide['diagram']:
            export_diagram(slide['diagram'], output/'diagrams')
    prs.save(ppt_path)
    canvas.save()
    sizes = [i['size'] for scene in scenes for i in scene
             if i['kind']=='text' and 160 < i['y'] < 490]
    minimum = min(sizes) if sizes else 0
    count = sum(bool(s['diagram']) for s in slides)
    image_count = sum(bool(s['image']) for s in slides)
    print(f'Rendered {len(slides)} slides / {count} flowcharts / {image_count} image pages; '
          f'minimum native content font {minimum:g} pt (excludes images)')
    print(ppt_path)
    print(pdf_path)


if __name__ == '__main__':
    main()
