"""A separate, native-object journey edition. Never writes the original deck or diagrams.

Run: .venv-slides/Scripts/python.exe tools/render_journey_slides.py
The current six-page edition preserves a read-only snapshot of the user's native
PPTX objects on pages 1–5, updates notes from Markdown, and redraws only the final
itinerary. The old nine-page builders remain available internally for history.
"""
from pathlib import Path
import argparse
import math
import os
import re
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader

import render_slides as base

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/presentations/journey/windows-cd-journey.md'
OUTPUT_NAME = 'six-page-v3'
MANUAL_TEMPLATE = SOURCE.parent / 'inputs/windows-cd-journey-manual-six.pptx'
W, H = 960, 540
INK = '#18394B'
MUTED = '#506B77'
PAPER = '#FCFBF7'
WHITE = '#FFFFFF'
BLUE = '#1676C4'
TEAL = '#00877F'
ORANGE = '#C86420'
RED = '#B33F3F'
LINE = '#DCE4E4'
SKY = '#EAF3FC'
MINT = '#E6F4ED'
SAND = '#FFF0DB'
ROSE = '#FAECE7'
COLORS = [BLUE, TEAL, ORANGE, BLUE, TEAL, ORANGE]
# Keep routes bright while darkening accent text for small-label contrast.
ACCENT_INK = {BLUE: '#126DB5', TEAL: '#00786F', ORANGE: '#AC4E16'}


class Scene:
    def __init__(self):
        self.items = []

    def box(self, x, y, w, h, fill=WHITE, stroke=None, form='round', radius=12):
        self.items.append(dict(kind='box', x=x, y=y, w=w, h=h, fill=fill,
                               stroke=stroke, form=form, radius=radius))

    def text(self, value, x, y, w, size=18, color=INK, bold=False, align='left', role=None):
        color = ACCENT_INK.get(color, color)
        lines = base.wrapped(value, w, size, bold)
        h = len(lines) * size * 1.35
        if x < 0 or y < 0 or x+w > W+.1 or y+h > H+.1:
            raise ValueError(f'Text outside slide: {value}')
        self.items.append(dict(kind='text', value=value, lines=lines, x=x, y=y,
                               w=w, h=h, size=size, color=color, bold=bold,
                               align=align, role=role))
        return h

    def path(self, points, color=INK, width=2, dashed=False, arrow=False, role=None):
        self.items.append(dict(kind='path', points=points, color=color,
                               width=width, dashed=dashed, arrow=arrow, role=role))

    def circle(self, x, y, r, fill=WHITE, stroke=None):
        self.box(x-r, y-r, 2*r, 2*r, fill, stroke, 'ellipse')

    def image(self, path, alt, title):
        iw, ih = ImageReader(str(path)).getSize()
        scale = min(W/iw, H/ih)
        self.items.append(dict(kind='image', path=path, alt=alt, title=title,
                               x=(W-iw*scale)/2, y=(H-ih*scale)/2,
                               w=iw*scale, h=ih*scale))


def bezier(a, b, c, d, steps=24):
    return [((1-t)**3*a[0]+3*(1-t)**2*t*b[0]+3*(1-t)*t*t*c[0]+t**3*d[0],
             (1-t)**3*a[1]+3*(1-t)**2*t*b[1]+3*(1-t)*t*t*c[1]+t**3*d[1])
            for t in (i/steps for i in range(steps+1))]


def route(s, stops, color=BLUE, width=12, track=True):
    points = []
    for a, d in zip(stops, stops[1:]):
        dx = (d[0]-a[0])*.48
        points += bezier(a, (a[0]+dx, a[1]), (d[0]-dx, d[1]), d)[bool(points):]
    s.path(points, color, width)
    if track:
        s.path(points, WHITE, 1.1)


def pill(s, label, x, y, w, color=TEAL, fill=MINT):
    s.box(x, y, w, 27, fill, radius=13)
    s.text(label, x+9, y+5, w-18, 11.5, color, True, 'center')


def node(s, x, y, label, color=BLUE, r=13):
    s.circle(x, y+3, r, '#DDE4E1')
    s.circle(x, y, r, WHITE, color)
    s.circle(x, y, r-4, ACCENT_INK.get(color, color))
    if label:
        s.text(label, x-r, y-7, r*2, 10, WHITE, True, 'center')


def icon(s, name, x, y, size=42, color=BLUE, fill=SKY):
    """Original geometric line illustrations; coordinates normalized to a 48-unit tile."""
    k = size/48
    def line(pts, width=2):
        s.path([(x+a*k, y+b*k) for a,b in pts], color, width*k)
    def box(a,b,w,h,form='round',f=fill):
        s.box(x+a*k,y+b*k,w*k,h*k,f,color,form,4*k)
    def dot(a,b,r=2,f=color):
        s.circle(x+a*k,y+b*k,r*k,f)
    if name == 'code':
        box(3,6,42,32)
        line([(3,14),(45,14)])
        for a in [8,13,18]: dot(a,10,1)
        line([(18,20),(12,25),(18,30)])
        line([(30,20),(36,25),(30,30)])
        line([(26,18),(22,33)])
        line([(17,43),(31,43)])
    elif name == 'gate':
        box(6,24,10,21)
        box(7,19,8,6,f=SAND)
        line([(11,24),(39,8)],5)
        line([(21,16),(24,21)],2.5)
        line([(30,11),(33,16)],2.5)
        dot(11,29,2,WHITE)
        line([(32,29),(37,34),(45,24)],2.5)
    elif name == 'server':
        for b in [5,18,31]:
            box(8,b,32,10)
            dot(13,b+5,1.5)
            line([(26,b+5),(35,b+5)],1.5)
    elif name == 'flag':
        line([(10,43),(10,5)],2.5)
        line([(10,6),(26,6),(26,12),(40,12),(34,22),(22,22),(22,16),(10,16)])
        s.box(x+11*k,y+7*k,14*k,8*k,fill,None,'rect')
        line([(4,43),(18,43)])
    elif name == 'check':
        box(10,8,30,35)
        box(17,4,15,8,f=WHITE)
        line([(16,21),(20,25),(28,17)])
        line([(16,33),(20,37),(28,29)])
        line([(31,23),(35,23)],1.5)
        line([(31,35),(35,35)],1.5)
    elif name == 'shield':
        line([(24,4),(42,11),(39,30),(33,39),(24,45),(15,39),(9,30),(6,11),(24,4)])
        line([(15,24),(22,31),(34,17)],3)
    elif name == 'branch':
        line([(12,8),(12,40)])
        line([(12,28),(28,28),(36,18),(36,8)])
        for a,b in [(12,8),(12,40),(36,8)]:
            s.circle(x+a*k,y+b*k,4*k,fill,color)
    elif name == 'clock':
        s.circle(x+24*k,y+26*k,18*k,fill,color)
        line([(24,14),(24,26),(34,26)])
        line([(19,3),(29,3)],3)
        line([(24,3),(24,7)])
    elif name == 'restore':
        line([(10,17),(16,9),(31,8),(40,18),(40,31),(31,40),(18,40),(9,32)],2.5)
        line([(10,5),(10,17),(22,17)],2.5)
        line([(24,18),(24,27),(31,31)])
    elif name == 'people':
        for a,b in [(15,14),(33,14),(24,8)]:
            s.circle(x+a*k,y+b*k,5*k,fill,color)
        for a,b in [(5,27),(28,27),(14,23)]: box(a,b,15,17)
    elif name == 'train':
        box(8,4,32,35)
        box(13,10,22,13,f=WHITE)
        dot(16,30,2)
        dot(32,30,2)
        line([(14,39),(9,45)])
        line([(34,39),(39,45)])
        line([(13,42),(35,42)])
    elif name == 'pipeline':
        for a in (2,19,36):
            box(a,17,10,16)
        line([(12,25),(19,25)])
        line([(29,25),(36,25)])
        line([(15,22),(18,25),(15,28)])
        line([(32,22),(35,25),(32,28)])
        line([(7,11),(7,6),(41,6),(41,11)],1.5)
    elif name == 'verify':
        box(5,8,25,33)
        box(11,4,13,8,f=WHITE)
        line([(11,20),(15,24),(23,16)])
        line([(11,31),(15,35),(23,27)])
        s.circle(x+35*k,y+32*k,11*k,fill,color)
        line([(29,32),(33,36),(41,27)],2.5)
    elif name == 'expand':
        box(17,3,14,12)
        line([(24,15),(24,25)])
        line([(7,31),(7,25),(41,25),(41,31)])
        line([(24,25),(24,31)])
        for a in (1,18,35):
            box(a,32,12,12)
    elif name == 'city':
        for a,b,w,h in [(2,22,12,23),(18,5,14,40),(36,14,10,31)]:
            box(a,b,w,h)
            for by in range(int(b+5),int(b+h-3),7):
                line([(a+3,by),(a+w-3,by)],1)
        line([(0,46),(48,46)])
    else:
        raise ValueError(f'Unknown illustration: {name}')


def footer(s, i, count=9):
    s.path([(36,505),(924,505)],LINE,.8)
    s.text('JOURNEY EDITION  /  規劃提案 · 尚未完成實作／POC 驗證',36,516,790,9.5,MUTED)
    s.text(f'{i+1:02d} / {count:02d}',864,514,60,11,MUTED,True,'center')


def page(i, title, subtitle, chapter, count=9):
    s = Scene()
    s.box(0,0,W,H,PAPER,form='rect')
    s.text('WINDOWS  /  DEPLOYMENT JOURNEY',36,19,615,10.5,TEAL,True)
    pill(s,chapter,766,15,158,INK,'#EEF0E9')
    s.text(title,36,59,888,29,INK,True,role='title')
    s.text(subtitle,36,104,888,16.5,MUTED)
    footer(s,i,count)
    return s


def cover(slide):
    s = Scene()
    s.box(0,0,W,H,PAPER,form='rect')
    s.text('WINDOWS  /  DEPLOYMENT JOURNEY',40,28,700,11,TEAL,True)
    pill(s,'另一種設計提案',778,22,142,INK,'#EEF0E9')
    s.text('保留建置的效率，\n走向受控的正式部署。',40,107,590,38,INK,True,role='title')
    s.text(slide['title'],40,222,550,23,INK,True)
    s.text('GitLab 建置  →  Azure DevOps Server 管理放行',40,264,650,18,MUTED)
    s.text('正式換版，由授權部門掌握。',40,294,590,19,TEAL,True)
    # This route is a reading metaphor, not a network/topology diagram.
    route(s,[(52,406),(250,406),(420,372),(610,392),(745,318),(886,253)],BLUE,13)
    route(s,[(610,392),(745,318),(886,253)],TEAL,13)
    for x,y,n,c in [(102,406,'01',BLUE),(420,372,'02',ORANGE),(745,318,'03',TEAL)]: node(s,x,y,n,c,15)
    icon(s,'code',75,339,52,BLUE,SKY)
    icon(s,'gate',394,302,56,ORANGE,SAND)
    icon(s,'city',797,155,108,TEAL,MINT)
    s.text('建置候選版本',53,432,180,16,BLUE,True)
    s.text('選版與授權',374,404,180,16,ORANGE,True)
    s.text('受控換版',705,347,180,16,TEAL,True)
    pill(s,'Express 作為 POC 起點',40,470,198,BLUE,SKY)
    pill(s,'目標 12/31 · 待共同確認',252,470,225,ORANGE,SAND)
    footer(s,0)
    return s


def agenda(slide):
    s = page(1,'六個問題，一條清楚的導入路線',slide['takeaway'],'01  /  今天的路線')
    groups = [('先看價值與邊界','01  預期效益','02  正式主機存取邊界','shield',BLUE,SKY),
              ('再看設計與分工','03  需求狀態與目標架構','04  分支與部署策略','branch',TEAL,MINT),
              ('最後走向落地','05  授權換版與復原設計','06  POC 到正式上線','flag',ORANGE,SAND)]
    route(s,[(70,337),(315,337),(400,337),(605,337),(705,337),(894,337)],TEAL,8)
    for j,(head,a,b,im,c,tint) in enumerate(groups):
        x=40+j*301
        s.box(x,168,278,145,tint)
        icon(s,im,x+18,183,48,c,tint)
        s.text(head,x+18,247,245,21,INK,True)
        node(s,x+28,337,str(j+1),c)
        s.text(a,x+9,369,275,18,INK,True)
        s.text(b,x+9,410,275,18,MUTED)
    return s


def benefits(slide):
    s=page(2,'讓換版不再每次重新摸索',slide['takeaway'],'02  /  預期效益')
    specs=[('clock',BLUE,SKY,'維運更一致','一致流程，少搬檔、少漏步。','人員、時間、主機與版本\n可追查。'),
           ('shield',TEAL,MINT,'降低操作風險','版本與目標綁定放行。','備份與復原演練，\n降低失敗影響。'),
           ('city',ORANGE,SAND,'複用到其他專案','留下治理與操作範本。','各專案的服務與復原方式，\n仍需個別驗證。')]
    for j,(im,c,tint,head,a,b) in enumerate(specs):
        x=40+j*301
        s.box(x,169,278,277,WHITE,LINE)
        s.box(x+17,186,70,70,tint)
        icon(s,im,x+30,197,46,c,tint)
        s.text(f'0{j+1}',x+209,198,47,27,c,True)
        s.text(head,x+18,275,244,23,INK,True)
        y=320
        y+=s.text(a,x+18,y,244,18,INK)+12
        s.text(b,x+18,y,244,17,MUTED)
    s.text('預期改善，不是既有成果；POC／試辦後再以紀錄比較。',40,467,880,17,MUTED)
    return s


def boundary(slide):
    s=page(3,'只送檔，也已越過正式主機存取邊界',slide['takeaway'],'03  /  已確認的邊界')
    s.box(40,161,268,193,SKY)
    icon(s,'code',58,178,49,BLUE,SKY)
    s.text('開發／CI 區',124,188,172,22,BLUE,True)
    s.text('開發人員提交程式碼\nGitLab Runner 建置／測試',60,251,229,18,INK)
    s.box(650,161,270,193,MINT)
    icon(s,'server',673,177,49,TEAL,MINT)
    s.text('Production',731,187,174,22,TEAL,True)
    s.text('正式主機僅限\n指定部門存取',675,251,225,20,INK,True)
    s.path([(326,260),(630,260)],RED,3,True,True)
    s.circle(478,260,24,ROSE,RED)
    s.path([(468,250),(488,270)],RED,3)
    s.path([(488,250),(468,270)],RED,3)
    pill(s,'禁止直接送檔',391,200,176,RED,ROSE)
    s.text('送檔 ≠ 尚未存取',356,304,255,19,RED,True,'center')
    route(s,[(171,354),(171,401),(476,401),(785,401),(785,354)],TEAL,7,False)
    s.box(332,369,302,70,SAND)
    icon(s,'gate',346,380,45,ORANGE,SAND)
    s.text('授權流程：放行後交付與換版',403,387,219,16,INK,True)
    s.text('保留建置自動化；不以共用目錄、帳號或中繼腳本繞過邊界。',40,468,880,17,MUTED)
    return s


def status(slide):
    s=page(4,'先把「需求」與「做得到」分開',slide['takeaway'],'04  /  證據與狀態')
    specs=[(BLUE,SKY,'需求已確認','shield'),(TEAL,MINT,'做法規劃中','server'),(ORANGE,SAND,'POC 待驗證','check')]
    for j,((c,tint,tag,im),card) in enumerate(zip(specs,slide['cards'])):
        x=40+j*301
        s.box(x,166,278,302,WHITE,LINE)
        s.box(x,166,278,7,c,form='rect')
        icon(s,im,x+18,191,43,c,tint)
        pill(s,tag,x+104,194,156,c,tint)
        s.text(card['title'],x+18,254,244,22,INK,True)
        y=303
        for bullet in card['bullets']:
            y+=s.text(bullet,x+18,y,242,17,INK)+18
    return s


def architecture_link(s, points, color=BLUE, width=2, dashed=False, role=None):
    """Round schematic elbows without changing endpoint/control semantics."""
    rounded = [points[0]]
    for a, b, c in zip(points, points[1:], points[2:]):
        before, after = math.dist(a, b), math.dist(b, c)
        radius = min(12, before/2, after/2)
        if not radius:
            continue
        entry = tuple(b[i] + (a[i]-b[i])*radius/before for i in (0, 1))
        exit = tuple(b[i] + (c[i]-b[i])*radius/after for i in (0, 1))
        ctrl1 = tuple(entry[i] + (b[i]-entry[i])*2/3 for i in (0, 1))
        ctrl2 = tuple(exit[i] + (b[i]-exit[i])*2/3 for i in (0, 1))
        rounded.append(entry)
        rounded.extend(bezier(entry, ctrl1, ctrl2, exit, 12)[1:])
    rounded.append(points[-1])
    s.path(rounded, color, width, dashed, True, role)


def architecture(slide):
    s=page(5,'從程式碼出發，在放行後抵達正式環境',
           '端到端目標架構／待 POC：建置、版本管理與正式部署，各自守住責任邊界。',
           '05  /  目標架構')
    s.text('本頁看交付與權限分工；單機換版、人工功能確認與復原流程 → 第 8 頁。',
           40,135,880,12.5,MUTED)
    # Three responsibility zones, not a physical network or mandatory VM split.
    for x,w,tint in [(40,212,SKY),(276,360,'#F1F3ED'),(660,260,MINT)]:
        s.box(x,160,w,299,tint)
    icon(s,'code',53,175,33,BLUE,SKY)
    s.text('開發／CI 區',99,176,140,20,BLUE,True)
    icon(s,'shield',290,175,33,ORANGE,SAND)
    s.text('授權部門管理區',335,176,280,20,INK,True)
    s.text('Azure DevOps Server',292,211,326,17,INK,True)
    icon(s,'server',674,175,33,TEAL,MINT)
    s.text('Production',719,176,190,20,TEAL,True)

    # Connections are drawn before the component cards. Artifact download is a
    # logical data direction, not a claim that management logs into Production.
    architecture_link(s,[(146,269),(146,286)],BLUE,2,True,'control')
    architecture_link(s,[(146,347),(146,378)],BLUE,2,True,'control')
    architecture_link(s,[(236,312),(292,312)],BLUE,2,True,'notification')
    architecture_link(s,[(236,411),(292,411)],BLUE,4,False,'artifact-import')
    architecture_link(s,[(363,342),(363,378)],BLUE,2,True,'control')
    architecture_link(s,[(435,411),(455,411)],BLUE,2,True,'version-selection')
    architecture_link(s,[(538,343),(538,378)],ORANGE,2,True,'authorization')
    s.box(56,225,180,44,WHITE)
    s.text('開發人員｜提交程式碼',67,238,158,14,INK,True,'center')
    s.box(56,286,180,61,WHITE)
    s.text('GitLab CE',68,292,156,19,INK,True)
    s.text('版本控制／CI 編排',68,322,156,12.5,MUTED)
    s.box(56,378,180,65,WHITE)
    s.text('Runner 建置／測試',68,387,156,15,INK,True)
    s.text('產物留存於 GitLab',68,416,156,12.5,MUTED)

    s.box(292,277,143,65,WHITE)
    s.text('接收成功通知',303,287,121,16,INK,True)
    s.text('驗證後啟動匯入',303,318,121,12,MUTED)
    s.box(292,378,143,65,WHITE)
    s.text('下載並保存版本',301,387,125,15.5,INK,True)
    s.text('管理 Build Agent',303,416,121,12,MUTED)
    s.box(455,244,165,99,SAND)
    icon(s,'gate',467,252,33,ORANGE,SAND)
    s.text('部署人員',507,253,102,17,ORANGE,True)
    s.text('選版＋應用／目標',468,296,140,15,INK,True)
    s.text('確認授權放行',468,321,140,13,ORANGE,True)
    s.box(455,378,165,65,WHITE)
    s.text('放行才部署',468,386,139,18,ORANGE,True)
    s.text('指定版本／指定目標',468,416,139,12,INK)

    # The cards are a stack of Windows Targets. Each has its own local Agent;
    # this is neither a central deployment host nor host-to-host file transfer.
    s.box(704,295,212,137,'#CBE6DA',TEAL)
    s.box(696,303,212,137,'#D8EDE3',TEAL)
    s.box(688,311,212,129,WHITE,TEAL)
    s.text('Windows Targets｜01 … N',700,319,188,13.5,INK,True,'center')
    s.box(702,349,182,28,SKY)
    s.text('每台主機的部署 Agent',709,355,168,13,BLUE,True,'center')
    s.box(702,400,182,26,MINT)
    s.text('各主機的應用服務',709,404,168,13,TEAL,True,'center')
    architecture_link(s,[(793,377),(793,400)],TEAL,2,True,'local-execution')
    # Draw these last so they visibly reach the local Agent, not just the host border.
    architecture_link(s,[(620,392),(670,392),(670,359),(702,359)],
                      BLUE,2,True,'dispatch')
    architecture_link(s,[(363,443),(363,471),(646,471),(646,370),(702,370)],
                      BLUE,4,False,'authorized-artifact')
    s.box(688,215,214,43,WHITE)
    s.text('F5｜應用流量',700,222,190,16,TEAL,True,'center')
    s.text('Agent 主動連 ADO 取工作與版本',675,270,232,12.5,INK,True,'center')
    architecture_link(s,[(902,236),(918,236),(918,449),(793,449),(793,426)],
                      TEAL,3,False,'application-traffic')

    # Compact legend: source/target semantics are carried by line type, not color alone.
    s.path([(48,485),(80,485)],BLUE,2,True)
    s.text('通知／控制',88,477,87,11.5,MUTED)
    s.path([(190,485),(222,485)],BLUE,4)
    s.text('版本產物',230,477,80,11.5,MUTED)
    s.path([(320,485),(352,485)],TEAL,3)
    s.text('應用流量',360,477,80,11.5,MUTED)
    s.text('管理側不碰正式主機；放行前不得預拷。',455,480,465,12,INK,True)
    s.text('放行後取指定版本',443,449,180,11.5,BLUE,True,'center')
    return s


def branches(slide):
    s=page(6,'不同 branch，各走合適的部署路線','規劃示意／待 POC：各分支各自建置；正式部署仍須授權。','06  /  分支與環境')
    for x,w,label in [(52,254,'來源分支'),(345,169,'觸發與策略'),(556,242,'執行者'),(813,110,'目標環境')]:
        s.text(label,x,151,w,15,MUTED,True)
    for y,c,tint in [(183,BLUE,SKY),(270,TEAL,MINT),(357,ORANGE,SAND)]:
        s.box(40,y,880,73 if y<350 else 100,tint)
        route(s,[(68,y+53 if y<350 else y+79),(285,y+53 if y<350 else y+79),
                 (326,y+47 if y<350 else y+73),(868,y+47 if y<350 else y+73)],c,7)
    icon(s,'branch',52,191,27,BLUE,SKY)
    s.text('master',90,192,227,20,INK,True)
    s.text('符合條件自動觸發',338,193,200,17,BLUE,True)
    s.text('Azure DevOps Server + Agent',555,193,242,15,INK,True)
    pill(s,'SIT',819,188,89,BLUE,WHITE)
    icon(s,'branch',52,278,27,TEAL,MINT)
    s.text('pre-production',90,279,236,19,INK,True)
    s.text('符合條件自動觸發',338,280,200,17,TEAL,True)
    s.text('Azure DevOps Server + Agent',555,280,242,15,INK,True)
    pill(s,'UAT',819,275,89,TEAL,WHITE)
    icon(s,'branch',52,369,27,ORANGE,SAND)
    s.text('production / production-bq1\nproduction-bq2 / production-nh\n…（平行正式分支示例）',90,365,230,12.8,INK,True)
    icon(s,'gate',339,372,47,ORANGE,SAND)
    s.text('選版＋選目標\n授權放行',397,365,151,17,ORANGE,True)
    s.text('Azure DevOps Server + Agent\n依目標採不同換版策略',555,365,244,15,INK,True)
    pill(s,'Production',802,366,110,ORANGE,WHITE)
    s.text('不同分支可有不同 pipeline；不將 production 前綴自動視為已授權。',40,475,883,16.5,MUTED)
    return s


def lifecycle(slide):
    s=page(7,'放行後才送檔，流量恢復前再確認','規劃／待 POC：Agent、備份、人工等待與復原腳本均須實作及演練。','07  /  換版與復原')
    # Native vector blocks explicitly retain the second, mid-deployment human gate.
    steps=[('01','人員選版／目標','授權放行','gate',ORANGE,SAND),
           ('02','Agent 取檔驗證','下 F5，確認隔離','server',BLUE,SKY),
           ('03','停止、備份','安裝新版並重啟','restore',TEAL,MINT),
           ('04','健康檢查正常','人工確認功能','check',ORANGE,SAND),
           ('05','上 F5 並確認','換版完成、留紀錄','flag',TEAL,MINT)]
    route(s,[(67,273),(895,273)],TEAL,7)
    for j,(n,a,b,im,c,tint) in enumerate(steps):
        x=40+j*179
        s.box(x,165,164,98,tint)
        icon(s,im,x+12,180,36,c,tint)
        s.text(n,x+109,187,40,23,c,True)
        node(s,x+23,273,'',c,9)
        s.text(a,x+1,296,165,17,INK,True)
        s.text(b,x+1,325,165,15.5,MUTED)
    s.path([(689,347),(689,373),(507,373)],RED,2,True,True)
    s.text('換版／啟動／檢查／新版功能驗證失敗',287,354,394,12.5,RED,True)
    s.box(40,391,421,87,ROSE)
    icon(s,'restore',52,409,37,RED,ROSE)
    s.text('保持隔離 → 還原舊版 → 再檢查',102,400,346,17,RED,True)
    s.text('舊版健康且人工確認，才可上 F5。\n復原成功仍記為「換版失敗」。',102,429,346,14.5,INK)
    s.box(481,391,439,87,SAND)
    s.text('未確認、逾時或復原失敗：不恢復流量',497,402,406,17,ORANGE,True)
    s.text('停止後續目標並通知；主機失聯另需復原手段。',497,441,406,14.5,INK)
    s.path([(507,373),(251,373),(251,391)],RED,2,True,True)
    return s


def legacy_timeline(slide):
    s=page(8,'從第一個 POC，走向受控的正式上線','暫定里程碑順序；階段日期、資源與上線批准，待相關部門共同確認。','08  /  建議行程')
    stops=[(86,348),(238,281),(390,324),(542,235),(694,280),(856,199)]
    for j in range(5):
        route(s,stops[j:j+2],COLORS[j],12)
    specs=[('POC','驗證通知與權限','隔離評估環境','code',44,382,SKY),
           ('協調架構部門','安裝／配置 ADO','確認責任與資源','people',167,169,MINT),
           ('整合 GitLab','串接通知與產物','保留與稽核','branch',320,361,SAND),
           ('SIT','整合與例外測試','包含權限負向測試','check',472,136,SKY),
           ('UAT','驗收與復原演練','由授權人員驗收','restore',625,319,MINT),
           ('Production','目標 12/31','依程序批准後試辦','flag',779,237,SAND)]
    for j,((x,y),(head,a,b,im,lx,ly,tint)) in enumerate(zip(stops,specs)):
        c=COLORS[j]
        node(s,x,y,f'0{j+1}',c,14)
        # Illustrations occupy the opposite side of each label group.
        iy=y-66 if j%2==0 else y+26
        if j==5:
            icon(s,'city',x-36,128,65,c,tint)
            icon(s,'flag',x+38,151,29,c,tint)
        else:
            icon(s,im,x-22,iy,43,c,tint)
        s.text(head,lx,ly,151,18,c,True)
        s.text(a,lx,ly+29,155,15,INK,True)
        s.text(b,lx,ly+53,155,13.5,MUTED)
    s.box(40,464,880,30,MINT)
    s.text('逐關通過再上線：先證明「未放行不能送檔」，再完成驗收與復原演練。',52,470,855,15,TEAL,True)
    return s


BUILDERS=[cover,agenda,benefits,boundary,status,architecture,branches,lifecycle,legacy_timeline]


def make_scenes(slides, source=SOURCE):
    if len(slides)!=9 or len(slides[5]['cards'])!=3 or slides[5]['image']:
        raise ValueError('Journey layout requires nine slides with a native three-zone architecture')
    return [builder(slide) for builder,slide in zip(BUILDERS,slides)]


def ppt_scene(prs, scene, notes, title, existing_slide=None):
    slide=existing_slide if existing_slide is not None else prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text=title
    slide.name=title[:200]
    for item in scene.items:
        kind=item['kind']
        if kind=='text':
            if item['role']=='title':
                box=slide.shapes.title
                box.left,box.top=Pt(item['x']),Pt(item['y'])
                box.width,box.height=Pt(item['w']),Pt(item['h']+2)
                parent=box._element.getparent()
                parent.remove(box._element)
                parent.append(box._element)
            else:
                box=slide.shapes.add_textbox(Pt(item['x']),Pt(item['y']),Pt(item['w']),Pt(item['h']+2))
            box.name='Title' if item['role']=='title' else 'Text'
            tf=box.text_frame
            tf.clear()
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.word_wrap=False
            tf.vertical_anchor=MSO_ANCHOR.TOP
            for i,line in enumerate(item['lines']):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                p.text=line
                p.alignment=PP_ALIGN.CENTER if item['align']=='center' else PP_ALIGN.LEFT
                p.space_before=p.space_after=Pt(0)
                p.line_spacing=Pt(item['size']*1.35)
                p.font.name=base.FONT
                p.font.size=Pt(item['size'])
                p.font.bold=item['bold']
                p.font.color.rgb=base.rgb(item['color'])
        elif kind=='image':
            pic=slide.shapes.add_picture(str(item['path']),Pt(item['x']),Pt(item['y']),width=Pt(item['w']),height=Pt(item['h']))
            pic._element.nvPicPr.cNvPr.set('descr',item['alt'])
        elif kind=='box':
            form={'round':MSO_SHAPE.ROUNDED_RECTANGLE,'ellipse':MSO_SHAPE.OVAL,'rect':MSO_SHAPE.RECTANGLE}[item['form']]
            shape=slide.shapes.add_shape(form,Pt(item['x']),Pt(item['y']),Pt(item['w']),Pt(item['h']))
            if item['form']=='round':
                shape.adjustments[0]=min(.5,item['radius']/min(item['w'],item['h']))
            shape.fill.solid()
            shape.fill.fore_color.rgb=base.rgb(item['fill'])
            if item['stroke']:
                shape.line.color.rgb=base.rgb(item['stroke'])
                shape.line.width=Pt(1.2)
            else: shape.line.fill.background()
        elif kind=='path':
            # Freeform polylines remain a single editable native PPTX object.
            # Coordinates use 1/100 point precision; bezier samples have no bitmap dependency.
            points=[(round(x*100),round(y*100)) for x,y in item['points']]
            builder=slide.shapes.build_freeform(*points[0],scale=127)
            builder.add_line_segments(points[1:],close=False)
            shape=builder.convert_to_shape()
            shape.fill.background()
            shape.line.color.rgb=base.rgb(item['color'])
            shape.line.width=Pt(item['width'])
            ln=shape.line._get_or_add_ln()
            ln.set('cap','rnd')
            join=base.OxmlElement('a:round')
            ln.append(join)
            if item['dashed']:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            if item['arrow']:
                end=base.OxmlElement('a:tailEnd')
                end.set('type','triangle')
                ln.append(end)
    slide.notes_slide.notes_text_frame.text=notes


def pdf_scene(canvas,scene):
    for i in scene.items:
        kind=i['kind']
        if kind=='text':
            canvas.setFont('bold' if i['bold'] else 'regular',i['size'])
            canvas.setFillColor(i['color'])
            for j,line in enumerate(i['lines']):
                x=i['x']
                if i['align']=='center':
                    x+=(i['w']-base.pdfmetrics.stringWidth(line,'bold' if i['bold'] else 'regular',i['size']))/2
                canvas.drawString(x,H-i['y']-i['size']*(1.08+j*1.35),line)
        elif kind=='box':
            canvas.setFillColor(i['fill'])
            canvas.setStrokeColor(i['stroke'] or i['fill'])
            canvas.setLineWidth(1.2)
            x,y,w,h=i['x'],H-i['y']-i['h'],i['w'],i['h']
            stroke=bool(i['stroke'])
            if i['form']=='ellipse': canvas.ellipse(x,y,x+w,y+h,fill=1,stroke=stroke)
            elif i['form']=='round': canvas.roundRect(x,y,w,h,min(i['radius'],min(w,h)/2),fill=1,stroke=stroke)
            else: canvas.rect(x,y,w,h,fill=1,stroke=stroke)
        elif kind=='path':
            canvas.setStrokeColor(i['color'])
            canvas.setLineWidth(i['width'])
            canvas.setLineCap(1)
            canvas.setLineJoin(1)
            canvas.setDash(5,4) if i['dashed'] else canvas.setDash()
            path=canvas.beginPath()
            path.moveTo(i['points'][0][0],H-i['points'][0][1])
            for x,y in i['points'][1:]: path.lineTo(x,H-y)
            canvas.drawPath(path)
            canvas.setDash()
            if i['arrow']:
                a,b=i['points'][-2:]
                angle=math.atan2(b[1]-a[1],b[0]-a[0])
                p=canvas.beginPath()
                p.moveTo(b[0],H-b[1])
                for side in (-1,1):
                    theta=angle+side*math.pi/6
                    p.lineTo(b[0]-8*math.cos(theta),H-b[1]+8*math.sin(theta))
                p.close()
                canvas.setFillColor(i['color'])
                canvas.drawPath(p,fill=1,stroke=0)
        elif kind=='image':
            canvas.drawImage(str(i['path']),i['x'],H-i['y']-i['h'],width=i['w'],height=i['h'])
    canvas.showPage()


def itinerary_entries(slide):
    """The six station labels and descriptions are sourced from Markdown, not code."""
    if len(slide['cards'])!=1 or len(slide['cards'][0]['bullets'])!=6:
        raise ValueError('The itinerary requires six station bullets separated by ｜')
    entries=[bullet.split('｜') for bullet in slide['cards'][0]['bullets']]
    if any(not 2<=len(parts)<=3 or any(not v.strip() for v in parts) for parts in entries):
        raise ValueError('Each station needs a heading and one or two descriptions')
    return entries


def timeline(slide):
    s=page(5,slide['title'],slide['takeaway'],'05  /  預計行程',count=6)
    stops=[(86,348),(238,281),(390,324),(542,235),(694,280),(856,199)]
    colors=[BLUE,TEAL,BLUE,TEAL,ORANGE,TEAL]
    tints=[SKY,MINT,SKY,MINT,SAND,MINT]
    symbols=['people','server','pipeline','verify','flag','expand']
    # Retain the user's rising route. The final stretch is a future extension,
    # not a second Production launch and not part of a new date commitment.
    for j in range(4):
        route(s,stops[j:j+2],colors[j],12)
    a,d=stops[4:]
    points=bezier(a,(a[0]+78,a[1]),(d[0]-78,d[1]),d)
    s.path(points,TEAL,7,dashed=True,role='future-expansion')
    label_positions=[(40,382,220),(164,182,242),(315,361,290),
                     (470,149,242),(625,319,293),(775,235,174)]
    illustration_positions=[(65,283,43),(217,309,43),(368,254,45),
                            (520,263,45),(674,210,43),(826,134,60)]
    for j,((x,y),parts) in enumerate(zip(stops,itinerary_entries(slide))):
        c=colors[j]
        node(s,x,y,f'0{j+1}',c,14)
        ix,iy,size=illustration_positions[j]
        icon(s,symbols[j],ix,iy,size,c,tints[j])
        lx,ly,width=label_positions[j]
        heading_size=16 if j!=2 else 17
        heading_height=s.text(parts[0],lx,ly,width,heading_size,c,True)
        body_y=ly+heading_height+6
        for k,line in enumerate(parts[1:]):
            body_y+=s.text(line,lx,body_y,width,13.5 if k==0 else 12.5,
                          INK if k==0 else MUTED,bold=k==0)+5
    # Replace the now-empty banner with an explicit scope distinction. Do not
    # imply that every future project must be delivered by the current target.
    s.box(40,470,880,27,'#F0F3EC')
    s.circle(57,483,4,ORANGE)
    s.text('本輪預計 12/28 前完成',70,475,270,12,ORANGE,True)
    s.path([(352,477),(352,490)],LINE,1)
    s.text('後續發展：逐案擴展，時程另議',374,475,515,12,TEAL,True)
    return s


def normalized(value):
    return re.sub(r'\s+','',value)


def validate_manual_template(prs,slides):
    if len(prs.slides)!=6 or len(slides)!=6:
        raise ValueError('Current edition requires six source and template pages')
    if (prs.slide_width,prs.slide_height)!=(Pt(W),Pt(H)):
        raise ValueError('Manual template must use the 960×540 layout')
    for n,(page_data,slide) in enumerate(zip(slides[:5],prs.slides)):
        actual=normalized('\n'.join(s.text for s in slide.shapes if s.has_text_frame))
        expected=[page_data['title'],page_data['takeaway']]
        for card in page_data['cards']:
            expected.extend([card['title'],*card['bullets']])
        for label in expected:
            if normalized(label) not in actual:
                raise ValueError(f'Page {n+1} content differs from manual template: {label}')
    itinerary_entries(slides[5])


def export_powerpoint_pdf(ppt_path,pdf_path):
    executable=shutil.which('powershell.exe')
    if not executable:
        raise RuntimeError('PowerPoint/Windows required for faithful PDF export; use --pptx-only')
    subprocess.run([executable,'-NoProfile','-ExecutionPolicy','Bypass','-File',
                    str(ROOT/'tools/export_presentation_pdf.ps1'),
                    '-InputPath',str(ppt_path.resolve()),'-OutputPath',str(pdf_path.resolve())],
                   check=True,timeout=120)
    if not pdf_path.is_file():
        raise RuntimeError('PowerPoint did not produce the PDF')


def build_six(slides,source,output_dir,template,export_pdf):
    prs=Presentation(template)
    validate_manual_template(prs,slides)
    scene=timeline(slides[5])
    # Preserve all the user's native shapes and manual formatting on pages 1–5.
    # Only notes and the nonvisible slide names are synchronized there.
    for slide,content in zip(prs.slides,slides):
        slide.notes_slide.notes_text_frame.text=content['notes']
        slide.name=content['title'][:200]
    last_slide=prs.slides[-1]
    title_shape=last_slide.shapes.title
    for shape in list(last_slide.shapes):
        if shape!=title_shape:
            shape._element.getparent().remove(shape._element)
    ppt_scene(prs,scene,slides[5]['notes'],slides[5]['title'],existing_slide=last_slide)
    prs.core_properties.title='Windows 自動化部署實作計畫｜六頁同步版'
    prs.core_properties.subject='依手改六頁版同步；12/28 為預計目標，尚待 POC 驗證與正式批准'
    output=output_dir or source.parent/OUTPUT_NAME
    if output.resolve()==template.parent.resolve():
        raise ValueError('Output must not overwrite the manual input directory')
    output.mkdir(parents=True,exist_ok=True)
    ppt_path=output/source.with_suffix('.pptx').name
    # Save in the destination filesystem, then atomically publish. A locked file
    # causes an error, never an automatic overwrite/fallback into another deck.
    with tempfile.TemporaryDirectory(dir=output) as staging:
        staged=Path(staging)/ppt_path.name
        prs.save(staged)
        os.replace(staged,ppt_path)
    if export_pdf:
        pdf_path=ppt_path.with_suffix('.pdf')
        with tempfile.TemporaryDirectory(dir=output) as staging:
            staged_pdf=Path(staging)/pdf_path.name
            export_powerpoint_pdf(ppt_path,staged_pdf)
            os.replace(staged_pdf,pdf_path)
    print(f'Created 6 slides; manual pages 1–5 retained, itinerary refreshed: {ppt_path}')
    return scene


def build(source=SOURCE,font_dir=None,output_dir=None,template=MANUAL_TEMPLATE,export_pdf=True):
    if source.stem != 'windows-cd-journey':
        raise ValueError('Use a windows-cd-journey.md copy; refusing to overwrite another deck')
    base.register_fonts(font_dir or Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts')
    slides=base.parse(source.read_text(encoding='utf-8'))
    if len(slides)==6:
        return build_six(slides,source,output_dir,template,export_pdf)
    scenes=make_scenes(slides,source)
    prs=Presentation()
    prs.slide_width,prs.slide_height=Pt(W),Pt(H)
    prs.core_properties.title='Windows 正式部署治理｜Journey 設計參考版'
    prs.core_properties.subject='獨立設計參考版；規劃提案，尚未實作或 POC 驗證'
    prs.core_properties.author='Windows CD Governance Project'
    output = output_dir or source.parent / OUTPUT_NAME
    output.mkdir(parents=True, exist_ok=True)
    ppt_path=output / source.with_suffix('.pptx').name
    pdf_path=output / source.with_suffix('.pdf').name
    canvas=Canvas(str(pdf_path),pagesize=(W,H),pageCompression=1)
    canvas.setTitle(prs.core_properties.title)
    for slide,scene in zip(slides,scenes):
        ppt_scene(prs,scene,slide['notes'],slide['title'])
        pdf_scene(canvas,scene)
    prs.save(ppt_path)
    canvas.save()
    print(f'Created {len(scenes)} slides; original deck and diagram assets untouched.')
    print(ppt_path)
    print(pdf_path)
    return scenes


def previews(source=SOURCE,output_dir=None):
    """Optional fixed-layout PDF previews; never represents a desktop PPTX playback test."""
    import pymupdf
    from PIL import Image, ImageDraw
    output = output_dir or source.parent / OUTPUT_NAME
    with pymupdf.open(output / source.with_suffix('.pdf').name) as document:
        rows = math.ceil(len(document)/3)
        sheet = Image.new('RGB', (1440, rows*285), LINE)
        draw = ImageDraw.Draw(sheet)
        for n, page in enumerate(document):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(.5, .5))
            image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            x, y = (n % 3)*480, (n // 3)*285
            sheet.paste(image, (x, y))
            draw.text((x+8, y+270), str(n+1), fill=INK)
            page.get_pixmap(matrix=pymupdf.Matrix(1.6, 1.6)).save(output/f'preview-{n+1:02d}.png')
        sheet.save(output/'overview.png')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=SOURCE)
    parser.add_argument('--font-dir',type=Path)
    parser.add_argument('--output-dir',type=Path,help='Default: six-page-v3 beside the source')
    parser.add_argument('--template',type=Path,default=MANUAL_TEMPLATE,help='Read-only manual six-page snapshot')
    parser.add_argument('--pptx-only',action='store_true',help='Skip PowerPoint PDF export')
    parser.add_argument('--previews',action='store_true',help='Also create PNG previews (requires PyMuPDF)')
    args=parser.parse_args()
    if args.pptx_only and args.previews:
        parser.error('--previews requires PDF export; omit --pptx-only')
    build(args.source,args.font_dir,args.output_dir,args.template,not args.pptx_only)
    if args.previews:
        previews(args.source,args.output_dir)
