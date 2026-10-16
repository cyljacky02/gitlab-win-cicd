"""Current six-page Journey edition: manual shapes are the editable layout source."""
from hashlib import sha256
from pathlib import Path
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.colors import HexColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import render_slides as base
import render_journey_slides as journey

FONT_DIR = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'


@unittest.skipUnless((FONT_DIR / 'msjh.ttc').is_file(), 'Microsoft JhengHei required')
class JourneyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        base.register_fonts(FONT_DIR)
        cls.pages = base.parse(journey.SOURCE.read_text(encoding='utf-8'))
        cls.template = Presentation(journey.MANUAL_TEMPLATE)
        cls.scene = journey.timeline(cls.pages[-1])

    def test_six_page_source_matches_manual_deck(self):
        self.assertEqual(len(self.pages), 6)
        journey.validate_manual_template(self.template, self.pages)
        self.assertEqual([p['title'] for p in self.pages], [
            '從程式建置到受控部署 建立可追蹤的換版流程',
            '讓換版不再每次重新摸索',
            '從程式碼出發，在放行後抵達正式環境',
            '不同 branch，各走合適的部署路線',
            '放行後才送檔，流量恢復前再確認',
            '從第一個 POC，走向受控的正式上線'])
        self.assertIn('Windows 自動化部署實作計畫', self.pages[0]['cards'][0]['title'])
        self.assertIn('第 5 頁', self.pages[2]['notes'])
        self.assertNotIn('第 8 頁', '\n'.join(p['notes'] for p in self.pages))
        self.assertNotIn('第 9 頁', '\n'.join(p['notes'] for p in self.pages))
        self.assertIn('只送檔也算正式環境存取', self.pages[2]['notes'])
        self.assertIn('不能代替中途人工功能關卡', self.pages[4]['notes'])
        self.assertTrue(all(p['notes'] for p in self.pages))

    def test_source_change_not_in_manual_template_fails_instead_of_silently_ignoring(self):
        from copy import deepcopy
        pages = deepcopy(self.pages)
        pages[0]['cards'][0]['bullets'][-1] = '預計 12/31 前完成'
        with self.assertRaisesRegex(ValueError, 'content differs'):
            journey.validate_manual_template(self.template, pages)

    def test_itinerary_uses_updated_manual_stages(self):
        entries = journey.itinerary_entries(self.pages[-1])
        self.assertEqual([e[0] for e in entries], [
            'POC & 跨部門討論', '協調架構部門', '串接 CI/CD Pipeline',
            'SIT & UAT 驗測', 'Production 正式環境', '後續發展'])
        self.assertEqual(entries[0][1], '(10/12 ~ 10/23)')
        self.assertIn('雷影', entries[2][1])
        visible = '\n'.join(i['value'] for i in self.scene.items if i['kind'] == 'text')
        for entry in entries:
            for part in entry:
                self.assertIn(part, visible)
        self.assertIn('12/28', visible)
        self.assertNotIn('12/31', visible)
        self.assertIn('時程另議', visible)
        self.assertIn('06 / 06', visible)
        future = [i for i in self.scene.items if i.get('role') == 'future-expansion']
        self.assertEqual(len(future), 1)
        self.assertTrue(future[0]['dashed'])
        self.assertEqual(future[0]['color'], journey.TEAL)

    def test_text_rectangles_do_not_overlap_or_cross_footer(self):
        texts = [i for i in self.scene.items if i['kind'] == 'text']
        for k, a in enumerate(texts):
            if a['y'] < 505:
                self.assertLessEqual(a['y'] + a['h'], 503, a['value'])
            for b in texts[k+1:]:
                dx = min(a['x']+a['w'], b['x']+b['w']) - max(a['x'], b['x'])
                dy = min(a['y']+a['h'], b['y']+b['h']) - max(a['y'], b['y'])
                self.assertFalse(dx > 1 and dy > 1, (a['value'], b['value']))

    def test_all_itinerary_geometry_inside_slide(self):
        for item in self.scene.items:
            if item['kind'] == 'path':
                points = item['points']
            else:
                points = [(item['x'], item['y']),
                          (item['x']+item['w'], item['y']+item['h'])]
            for x, y in points:
                self.assertTrue(0 <= x <= journey.W and 0 <= y <= journey.H, item)

    def test_accent_text_contrast_on_tinted_background(self):
        def luminance(color):
            return sum(w*(v/12.92 if v <= .04045 else ((v+.055)/1.055)**2.4)
                       for w, v in zip((.2126, .7152, .0722), HexColor(color).rgb()))
        for color, tint in [(journey.BLUE, journey.SKY), (journey.TEAL, journey.MINT),
                            (journey.ORANGE, journey.SAND)]:
            ratio = (luminance(tint)+.05)/(luminance(journey.ACCENT_INK[color])+.05)
            self.assertGreaterEqual(ratio, 4.5)

    def test_native_pages_one_to_five_preserved_and_all_notes_updated(self):
        protected = [journey.ROOT / p for p in (
            'docs/presentations/windows-cd-initial-plan.md',
            'docs/presentations/windows-cd-initial-plan.pptx',
            'docs/presentations/windows-cd-initial-plan.pdf',
            'docs/presentations/README.md', 'tools/render_slides.py',
            'tests/test_slide_images.py')]
        for directory in ('architecture-v2', 'inputs'):
            protected += list((journey.SOURCE.parent / directory).glob('*'))
        protected += [journey.SOURCE.parent / name for name in
                      ('windows-cd-journey.pptx', 'windows-cd-journey.pdf', 'overview.png')]
        protected += list((journey.ROOT / 'docs/presentations/diagrams').glob('*'))
        fingerprints = {p: sha256(p.read_bytes()).digest() for p in protected
                        if p.is_file() and not p.name.startswith('~$')}
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'windows-cd-journey.md'
            source.write_text(journey.SOURCE.read_text(encoding='utf-8'), encoding='utf-8')
            journey.build(source, FONT_DIR, export_pdf=False)
            output = source.parent / journey.OUTPUT_NAME
            prs = Presentation(output / source.with_suffix('.pptx').name)
            self.assertEqual(len(prs.slides), 6)
            self.assertEqual(prs.core_properties.title, 'Windows 自動化部署實作計畫｜六頁同步版')
            for n, slide in enumerate(prs.slides):
                self.assertIsNotNone(slide.shapes.title)
                self.assertEqual(slide.notes_slide.notes_text_frame.text, self.pages[n]['notes'])
                self.assertEqual(slide.name, self.pages[n]['title'])
                self.assertFalse(any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes))
                if n < 5:
                    self.assertEqual(slide.shapes._spTree.xml,
                                     self.template.slides[n].shapes._spTree.xml,
                                     f'Manual shapes changed on page {n+1}')
            self.assertNotEqual(prs.slides[-1].shapes._spTree.xml,
                                self.template.slides[-1].shapes._spTree.xml)
            self.assertFalse((output / source.with_suffix('.pdf').name).exists())
        self.assertEqual(fingerprints, {p: sha256(p.read_bytes()).digest() for p in fingerprints})

    def test_pdf_failure_does_not_publish_an_incomplete_pdf(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp)
            with patch.object(journey, 'export_powerpoint_pdf', side_effect=RuntimeError('Export failed')):
                with self.assertRaisesRegex(RuntimeError, 'Export failed'):
                    journey.build(output_dir=output)
            self.assertTrue((output / 'windows-cd-journey.pptx').exists())
            self.assertFalse((output / 'windows-cd-journey.pdf').exists())

    def test_refuses_original_filename_and_manual_input_output_directory(self):
        with self.assertRaisesRegex(ValueError, 'refusing to overwrite'):
            journey.build(journey.ROOT / 'docs/presentations/windows-cd-initial-plan.md')
        with self.assertRaisesRegex(ValueError, 'manual input directory'):
            journey.build(output_dir=journey.MANUAL_TEMPLATE.parent, export_pdf=False)


if __name__ == '__main__':
    unittest.main()
