"""Run: .venv-slides/Scripts/python -m unittest discover -s tests"""
import io
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.pdfgen.canvas import Canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import render_slides as slides


class ImageSlideTests(unittest.TestCase):
    SOURCE = '# Architecture\n> Controlled release\n\n![Editable source is draw.io](map.png)\n\n<!-- notes\nDetails\n-->'

    def test_parse_image_and_notes(self):
        page, = slides.parse(self.SOURCE)
        self.assertEqual(page['image'], {'alt': 'Editable source is draw.io', 'path': 'map.png'})
        self.assertIsNone(page['diagram'])
        self.assertEqual(page['notes'], 'Details')
        self.assertEqual(page['cards'], [])

    def test_reject_mixed_and_multiple_images(self):
        for extra in ['\n## Card\n- Text', '\n![Second](two.png)']:
            with self.subTest(extra=extra), self.assertRaises(ValueError):
                slides.parse(self.SOURCE + extra)

    def test_fit_without_crop_and_export(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'map.png'
            Image.new('RGB', (1000, 1000), 'white').save(path)
            page, = slides.parse(self.SOURCE)
            page['image']['path'] = path
            scene = slides.make_scene(page, 3, 7)
            image, = scene
            self.assertEqual((image['x'], image['y'], image['w'], image['h']), (210, 0, 540, 540))
            prs = Presentation()
            slides.ppt_scene(prs, scene, page['notes'])
            slide = prs.slides[0]
            self.assertEqual(slide.shapes.title.text, 'Architecture')
            pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            self.assertEqual(pictures[0]._element.nvPicPr.cNvPr.get('descr'), 'Editable source is draw.io')
            self.assertEqual(slide.notes_slide.notes_text_frame.text, 'Details')
            prs.save(Path(directory) / 'test.pptx')
            stream = io.BytesIO()
            canvas = Canvas(stream, pagesize=(slides.W, slides.H))
            slides.pdf_scene(canvas, scene)
            canvas.save()
            self.assertTrue(stream.getvalue().startswith(b'%PDF'))

    def test_value_and_evidence_status_are_visible(self):
        source = Path(__file__).resolve().parents[1] / 'docs/presentations/windows-cd-initial-plan.md'
        pages = slides.parse(source.read_text(encoding='utf-8'))
        self.assertIn('預期效益', pages[2]['title'])
        self.assertEqual([c['title'] for c in pages[2]['cards']],
                         ['維運更一致', '降低操作風險', '複用到其他專案'])
        self.assertEqual([c['title'] for c in pages[4]['cards']],
                         ['已確認的需求', '規劃中的做法', 'POC 待驗證'])
        for index in (6, 7):
            self.assertIn('規劃', pages[index]['title'])
            self.assertIn('POC', pages[index]['takeaway'])

    def test_current_deck_has_one_image(self):
        source = Path(__file__).resolve().parents[1] / 'docs/presentations/windows-cd-initial-plan.md'
        pages = slides.parse(source.read_text(encoding='utf-8'))
        self.assertEqual(len(pages), 9)
        self.assertEqual([i for i, p in enumerate(pages) if p['image']], [5])
        self.assertEqual(sum(bool(p['diagram']) for p in pages), 4)
        self.assertTrue((source.parent / pages[5]['image']['path']).is_file())


if __name__ == '__main__':
    unittest.main()
